#!/usr/bin/env python3
"""Generate the LayerLoom tensile-screen figures and specimen summary table.

The input is the raw Mark-10 Excel export directory.  This is intentionally a
descriptive strength screen: crosshead strain is retained for the curve plot,
but no modulus or yield values are calculated from machine displacement.
"""

from __future__ import annotations

import argparse
import csv
import math
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrowPatch, Polygon, Rectangle
from openpyxl import load_workbook


SPEED_MM_PER_MIN = 5.0
FLAT_GAUGE_WIDTH_MM = 3.06
FLAT_GAUGE_THICKNESS_MM = 2.07
UPRIGHT_GAUGE_WIDTH_MM = 6.00
UPRIGHT_GAUGE_THICKNESS_MM = 4.13

MAGENTA = "#cf2f84"
CYAN = "#00a8c6"
INK = "#1d2730"
MUTED = "#667681"
GRID = "#dbe1e5"
LIGHT_MAGENTA = "#f4d7e6"
LIGHT_CYAN = "#d9f0f4"
MAGENTA_SHADES = ("#7f174a", "#a91f62", "#cf2f84", "#df69a5", "#ee9fc4")
CYAN_SHADES = ("#006778", "#00899f", "#00a8c6", "#45bfd2", "#8ad9e3")

T_CRITICAL_95 = {3: 4.30265272975, 5: 2.7764451052}


@dataclass(frozen=True)
class Specimen:
    specimen_id: str
    source_filename: str
    condition: str
    condition_label: str
    orientation: str
    orientation_label: str
    scale: str
    replicate: int
    is_primary: bool
    status: str
    width_mm: float
    thickness_mm: float
    grip_separation_mm: float
    time_s: tuple[float, ...]
    tensile_force_n: tuple[float, ...]
    curve_stop_index: int

    @property
    def area_mm2(self) -> float:
        return self.width_mm * self.thickness_mm

    @property
    def peak_index(self) -> int:
        return max(range(len(self.tensile_force_n)), key=self.tensile_force_n.__getitem__)

    @property
    def peak_force_n(self) -> float:
        return self.tensile_force_n[self.peak_index]

    @property
    def peak_time_s(self) -> float:
        return self.time_s[self.peak_index]

    @property
    def peak_uts_mpa(self) -> float:
        return self.peak_force_n / self.area_mm2

    @property
    def peak_travel_mm(self) -> float:
        return SPEED_MM_PER_MIN * self.peak_time_s / 60.0

    @property
    def peak_crosshead_strain(self) -> float:
        return self.peak_travel_mm / self.grip_separation_mm

    @property
    def release_time_s(self) -> float:
        return self.time_s[self.curve_stop_index]

    @property
    def release_travel_mm(self) -> float:
        return SPEED_MM_PER_MIN * self.release_time_s / 60.0

    @property
    def release_crosshead_strain(self) -> float:
        return self.release_travel_mm / self.grip_separation_mm

    @property
    def curve_time_s(self) -> tuple[float, ...]:
        return self.time_s[: self.curve_stop_index + 1]

    @property
    def curve_force_n(self) -> tuple[float, ...]:
        return self.tensile_force_n[: self.curve_stop_index + 1]

    def crosshead_work_density_mj_m3(self, stop_index: int) -> float:
        """Integrate engineering stress over crosshead strain through an index."""
        stop_index = min(stop_index, len(self.time_s) - 1)
        strain = [
            SPEED_MM_PER_MIN * time_s / 60.0 / self.grip_separation_mm
            for time_s in self.time_s[: stop_index + 1]
        ]
        stress = [
            max(0.0, force_n) / self.area_mm2
            for force_n in self.tensile_force_n[: stop_index + 1]
        ]
        # MPa integrated over unitless strain is numerically MJ/m^3.
        return sum(
            0.5 * (stress[index - 1] + stress[index]) * (strain[index] - strain[index - 1])
            for index in range(1, len(strain))
        )

    @property
    def crosshead_work_density_to_peak_mj_m3(self) -> float:
        return self.crosshead_work_density_mj_m3(self.peak_index)

    @property
    def crosshead_work_density_to_release_mj_m3(self) -> float:
        return self.crosshead_work_density_mj_m3(self.curve_stop_index)


def parse_specimen(path: Path) -> Specimen:
    """Read a raw Mark-10 export, normalized so tensile force is positive."""
    name = path.stem
    match = re.fullmatch(r"(pink|purple)_(laying|standing)_(\d+)(?:_example)?", name)
    if not match:
        raise ValueError(f"Unrecognized tensile export filename: {path.name}")

    color, position, replicate_text = match.groups()
    replicate = int(replicate_text)
    condition = "magenta" if color == "pink" else "cm_weave"
    condition_label = (
        "Magenta single-color control" if condition == "magenta" else "CM LayerLoom weave"
    )
    orientation = "flat" if position == "laying" else "upright"
    orientation_label = "Flat, half-scale" if orientation == "flat" else "Upright, full-scale"
    scale = "half" if orientation == "flat" else "full"
    width_mm, thickness_mm, grip_separation_mm = (
        (FLAT_GAUGE_WIDTH_MM, FLAT_GAUGE_THICKNESS_MM, 40.0)
        if orientation == "flat"
        else (UPRIGHT_GAUGE_WIDTH_MM, UPRIGHT_GAUGE_THICKNESS_MM, 90.0)
    )

    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = workbook["Sheet2"]
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    header = rows[0]
    load_index = header.index("Load [N]")
    time_index = header.index("Time [sec.]")
    points = [
        (float(row[time_index]), -float(row[load_index]))
        for row in rows[1:]
        if row[load_index] is not None and row[time_index] is not None
    ]
    if len(points) < 3:
        raise ValueError(f"Too few force-displacement points in {path.name}")

    time_s, tensile_force_n = map(tuple, zip(*points))
    peak_index = max(range(len(tensile_force_n)), key=tensile_force_n.__getitem__)
    peak_force = tensile_force_n[peak_index]
    # Plot only through the first clear post-peak load release.  This prevents
    # the machine return stroke from being interpreted as material response.
    release_threshold = max(10.0, peak_force * 0.10)
    curve_stop_index = len(tensile_force_n) - 1
    for index in range(peak_index + 1, len(tensile_force_n)):
        if tensile_force_n[index] <= release_threshold:
            curve_stop_index = index
            break

    is_primary = orientation == "flat" or replicate >= 3
    status = "Primary analysis"
    if not is_primary:
        status = "Chronological familiarization; excluded from primary analysis"
    if name == "purple_standing_1_example":
        status += "; sign reversal after peak, curve truncated at release"

    return Specimen(
        specimen_id=name,
        source_filename=path.name,
        condition=condition,
        condition_label=condition_label,
        orientation=orientation,
        orientation_label=orientation_label,
        scale=scale,
        replicate=replicate,
        is_primary=is_primary,
        status=status,
        width_mm=width_mm,
        thickness_mm=thickness_mm,
        grip_separation_mm=grip_separation_mm,
        time_s=time_s,
        tensile_force_n=tensile_force_n,
        curve_stop_index=curve_stop_index,
    )


def mean(values: Iterable[float]) -> float:
    values = tuple(values)
    return sum(values) / len(values)


def sample_sd(values: Iterable[float]) -> float:
    values = tuple(values)
    average = mean(values)
    return math.sqrt(sum((value - average) ** 2 for value in values) / (len(values) - 1))


def ci95(values: Iterable[float]) -> float:
    values = tuple(values)
    return T_CRITICAL_95[len(values)] * sample_sd(values) / math.sqrt(len(values))


def primary_group(specimens: Iterable[Specimen], orientation: str, condition: str) -> list[Specimen]:
    return [
        specimen
        for specimen in specimens
        if specimen.is_primary and specimen.orientation == orientation and specimen.condition == condition
    ]


def add_panel_label(axis: plt.Axes, letter: str) -> None:
    axis.text(
        -0.08,
        1.04,
        letter,
        transform=axis.transAxes,
        ha="left",
        va="top",
        fontsize=12,
        fontweight="bold",
        color=INK,
    )


def draw_dogbone(axis: plt.Axes, x: float, y: float, width: float, height: float, *, woven: bool) -> None:
    """Draw a compact dogbone schematic, with visible CM layers when woven."""
    half_wide = width / 2
    half_narrow = width * 0.28
    shoulder = height * 0.22
    points = [
        (x - half_wide, y - height / 2),
        (x + half_wide, y - height / 2),
        (x + half_wide, y - height / 2 + shoulder),
        (x + half_narrow, y - height / 2 + shoulder * 1.2),
        (x + half_narrow, y + height / 2 - shoulder * 1.2),
        (x + half_wide, y + height / 2 - shoulder),
        (x + half_wide, y + height / 2),
        (x - half_wide, y + height / 2),
        (x - half_wide, y + height / 2 - shoulder),
        (x - half_narrow, y + height / 2 - shoulder * 1.2),
        (x - half_narrow, y - height / 2 + shoulder * 1.2),
        (x - half_wide, y - height / 2 + shoulder),
    ]
    outer = Polygon(points, closed=True, facecolor=MAGENTA if not woven else LIGHT_CYAN, edgecolor=INK, lw=0.8)
    axis.add_patch(outer)
    if woven:
        clip = Polygon(points, closed=True, transform=axis.transData)
        for offset in range(-8, 10):
            band = Rectangle(
                (x - half_wide - 0.02, y - height / 2 + offset * height / 10),
                width + 0.04,
                height / 10,
                facecolor=CYAN if offset % 2 == 0 else MAGENTA,
                edgecolor="none",
            )
            band.set_clip_path(clip)
            axis.add_patch(band)
        axis.add_patch(Polygon(points, closed=True, fill=False, edgecolor=INK, lw=0.8))


def plot_test_architecture(axis: plt.Axes) -> None:
    axis.set_xlim(0, 1)
    axis.set_ylim(0, 1)
    axis.axis("off")
    add_panel_label(axis, "A")
    axis.text(0.02, 0.96, "Test architecture", fontsize=10, fontweight="bold", color=INK, va="top")
    axis.text(0.02, 0.88, "PLA; 100% infill", ha="left", va="top", fontsize=7.4, color=MUTED)
    axis.text(0.02, 0.83, "0.16 mm weave step; 5 mm/min", ha="left", va="top", fontsize=7.4, color=MUTED)

    draw_dogbone(axis, 0.15, 0.52, 0.18, 0.36, woven=False)
    draw_dogbone(axis, 0.46, 0.52, 0.18, 0.36, woven=True)
    axis.text(0.15, 0.30, "Magenta\ncontrol", ha="center", va="top", fontsize=6.9, color=INK)
    axis.text(0.46, 0.30, "CM LayerLoom\nweave", ha="center", va="top", fontsize=6.9, color=INK)

    axis.add_patch(FancyArrowPatch((0.095, 0.73), (0.095, 0.31), arrowstyle="<->", mutation_scale=10, lw=1, color=INK))
    axis.text(0.06, 0.52, "tension", rotation=90, ha="center", va="center", fontsize=6.7, color=INK)
    axis.text(0.02, 0.09, "Flat, half-scale\nIn-plane loading", ha="left", va="center", fontsize=8.0, fontweight="bold", color=INK)

    axis.plot([0.65, 0.96], [0.18, 0.18], color=INK, lw=0.8)
    for x in (0.70, 0.78, 0.86):
        axis.plot([x, x + 0.05], [0.18, 0.21], color=GRID, lw=0.7)
    draw_dogbone(axis, 0.80, 0.50, 0.14, 0.47, woven=True)
    axis.add_patch(FancyArrowPatch((0.80, 0.77), (0.80, 0.23), arrowstyle="<->", mutation_scale=10, lw=1, color=INK))
    axis.text(0.84, 0.50, "tension", rotation=90, ha="center", va="center", fontsize=6.7, color=INK)
    axis.text(0.64, 0.09, "Upright, full-scale\nThrough-layer loading", ha="left", va="center", fontsize=8.0, fontweight="bold", color=INK)


def plot_uts(axis: plt.Axes, specimens: list[Specimen]) -> None:
    add_panel_label(axis, "B")
    groups = [
        ("flat", "magenta", "Flat, half-scale", "Magenta\ncontrol"),
        ("flat", "cm_weave", "Flat, half-scale", "CM weave"),
        ("upright", "magenta", "Upright, full-scale", "Magenta\ncontrol"),
        ("upright", "cm_weave", "Upright, full-scale", "CM weave"),
    ]
    positions = [0.0, 0.9, 2.6, 3.5]
    jitters = [-0.075, -0.035, 0.0, 0.035, 0.075]
    for (orientation, condition, _, _), x, color in zip(groups, positions, [MAGENTA, CYAN, MAGENTA, CYAN]):
        group = primary_group(specimens, orientation, condition)
        values = [specimen.peak_uts_mpa for specimen in group]
        offsets = jitters[: len(values)]
        axis.scatter(
            [x + offset for offset in offsets],
            values,
            s=28,
            c=color,
            edgecolors=INK,
            linewidths=0.55,
            zorder=3,
        )
        group_mean = mean(values)
        interval = ci95(values)
        axis.errorbar(x, group_mean, yerr=interval, fmt="none", color=INK, lw=1.15, capsize=3, zorder=2)
        axis.plot([x - 0.13, x + 0.13], [group_mean, group_mean], color=INK, lw=2.0, zorder=4)
        axis.text(x, 4.0, f"n={len(values)}", ha="center", va="bottom", fontsize=7.5, color=MUTED)

    flat_magenta = mean(specimen.peak_uts_mpa for specimen in primary_group(specimens, "flat", "magenta"))
    flat_cm = mean(specimen.peak_uts_mpa for specimen in primary_group(specimens, "flat", "cm_weave"))
    upright_magenta = mean(specimen.peak_uts_mpa for specimen in primary_group(specimens, "upright", "magenta"))
    upright_cm = mean(specimen.peak_uts_mpa for specimen in primary_group(specimens, "upright", "cm_weave"))
    axis.text(0.45, 42.6, f"CM = {flat_cm / flat_magenta * 100:.1f}% of control", ha="center", va="top", fontsize=7.5, color=INK)
    axis.text(3.05, 42.6, f"CM = {upright_cm / upright_magenta * 100:.1f}% of control", ha="center", va="top", fontsize=7.5, color=INK)
    axis.axvline(1.75, color=GRID, lw=0.8, zorder=0)
    axis.set_xlim(-0.45, 3.95)
    axis.set_ylim(0, 44)
    axis.set_ylabel("Ultimate tensile strength (MPa)", fontsize=9)
    axis.set_xticks(positions)
    axis.set_xticklabels([group[3] for group in groups], fontsize=7.5)
    axis.text(0.45, -0.24, "Flat, half-scale", transform=axis.get_xaxis_transform(), ha="center", va="top", fontsize=8.5, fontweight="bold", color=INK)
    axis.text(3.05, -0.24, "Upright, full-scale", transform=axis.get_xaxis_transform(), ha="center", va="top", fontsize=8.5, fontweight="bold", color=INK)
    axis.grid(axis="y", color=GRID, lw=0.7)
    axis.spines[["top", "right"]].set_visible(False)
    axis.tick_params(axis="y", labelsize=8)


def plot_main_figure(specimens: list[Specimen], output_path: Path) -> None:
    figure = plt.figure(figsize=(7.35, 3.15), constrained_layout=True)
    grid = figure.add_gridspec(1, 2, width_ratios=[1.05, 1.35], wspace=0.21)
    plot_test_architecture(figure.add_subplot(grid[0, 0]))
    plot_uts(figure.add_subplot(grid[0, 1]), specimens)
    figure.savefig(output_path, dpi=600, bbox_inches="tight", facecolor="white")
    plt.close(figure)


def plot_crosshead_curves(specimens: list[Specimen], output_path: Path) -> None:
    """Separate individual traces so post-peak behavior remains legible."""
    figure, axes = plt.subplots(2, 2, figsize=(7.2, 5.45), constrained_layout=True)
    panels = (
        ("flat", "magenta", "A", "Flat, half-scale", "Magenta control"),
        ("flat", "cm_weave", "B", "Flat, half-scale", "CM LayerLoom weave"),
        ("upright", "magenta", "C", "Upright, full-scale", "Magenta control"),
        ("upright", "cm_weave", "D", "Upright, full-scale", "CM LayerLoom weave"),
    )
    for axis, (orientation, condition, letter, orientation_title, condition_title) in zip(
        axes.flat, panels
    ):
        add_panel_label(axis, letter)
        plotted = sorted(
            (
                specimen
                for specimen in specimens
                if specimen.orientation == orientation and specimen.condition == condition
            ),
            key=lambda specimen: specimen.replicate,
        )
        shades = MAGENTA_SHADES if condition == "magenta" else CYAN_SHADES
        for specimen, color in zip(plotted, shades):
            strain = [
                SPEED_MM_PER_MIN * time_s / 60.0 / specimen.grip_separation_mm * 100.0
                for time_s in specimen.curve_time_s
            ]
            # A negative endpoint can reflect frame reversal after fracture;
            # it is a release event, not compressive material response.
            stress = [
                max(0.0, force_n) / specimen.area_mm2
                for force_n in specimen.curve_force_n
            ]
            style = "-" if specimen.is_primary else ":"
            alpha = 0.94 if specimen.is_primary else 0.58
            label = f"Rep. {specimen.replicate}"
            if not specimen.is_primary:
                label += " (fam.)"
            axis.plot(
                strain,
                stress,
                color=color,
                lw=1.35,
                ls=style,
                alpha=alpha,
                label=label,
            )
            axis.scatter(
                [specimen.peak_crosshead_strain * 100.0],
                [specimen.peak_uts_mpa],
                s=17,
                facecolor="white",
                edgecolor=color,
                linewidth=0.85,
                zorder=4,
            )

        axis.set_title(
            f"{orientation_title}\n{condition_title}",
            fontsize=9.0,
            fontweight="bold",
            color=INK,
            pad=5,
        )
        axis.set_xlabel("Crosshead strain (%)", fontsize=8.2)
        axis.set_ylabel("Engineering stress (MPa)", fontsize=8.2)
        axis.grid(color=GRID, lw=0.65)
        axis.spines[["top", "right"]].set_visible(False)
        axis.tick_params(labelsize=7.5)
        axis.legend(
            loc="upper left",
            frameon=False,
            fontsize=6.0,
            ncol=2 if orientation == "upright" else 3,
            columnspacing=0.8,
            handlelength=1.5,
            handletextpad=0.35,
        )
        if orientation == "flat":
            axis.set_xlim(0, 26.2)
            axis.set_ylim(0, 39.2)
        else:
            axis.set_xlim(0, 4.75)
            axis.set_ylim(0, 18.5)

    figure.savefig(output_path, dpi=600, bbox_inches="tight", facecolor="white")
    plt.close(figure)


def plot_crosshead_endpoints(specimens: list[Specimen], output_path: Path) -> None:
    """Plot exploratory crosshead-derived endpoints for primary specimens."""
    figure, axes = plt.subplots(1, 2, figsize=(7.2, 3.25), constrained_layout=True)
    groups = [
        ("flat", "magenta", "Magenta\ncontrol"),
        ("flat", "cm_weave", "CM weave"),
        ("upright", "magenta", "Magenta\ncontrol"),
        ("upright", "cm_weave", "CM weave"),
    ]
    positions = [0.0, 0.9, 2.6, 3.5]
    colors = [MAGENTA, CYAN, MAGENTA, CYAN]
    metrics = (
        (
            "A",
            lambda specimen: specimen.peak_crosshead_strain * 100.0,
            r"Strain at UTS, $\varepsilon_{\mathrm{ch},p}$ (%)",
            (0, 7.2),
        ),
        (
            "B",
            lambda specimen: specimen.crosshead_work_density_to_peak_mj_m3,
            r"Work to UTS, $w_{\mathrm{ch},p}$ (MJ m$^{-3}$)",
            (0, 1.22),
        ),
    )
    jitters = (-0.075, -0.035, 0.0, 0.035, 0.075)
    for axis, (letter, metric, ylabel, ylim) in zip(axes, metrics):
        add_panel_label(axis, letter)
        for (orientation, condition, _), x, color in zip(groups, positions, colors):
            group = primary_group(specimens, orientation, condition)
            values = [metric(specimen) for specimen in group]
            axis.scatter(
                [x + offset for offset in jitters[: len(values)]],
                values,
                s=27,
                c=color,
                edgecolors=INK,
                linewidths=0.5,
                zorder=3,
            )
            group_mean = mean(values)
            interval = ci95(values)
            axis.errorbar(
                x,
                group_mean,
                yerr=interval,
                fmt="none",
                color=INK,
                lw=1.05,
                capsize=3,
                zorder=2,
            )
            axis.plot(
                [x - 0.13, x + 0.13],
                [group_mean, group_mean],
                color=INK,
                lw=1.8,
                zorder=4,
            )
            axis.text(
                x,
                ylim[0] + 0.035 * (ylim[1] - ylim[0]),
                f"n={len(values)}",
                ha="center",
                va="bottom",
                fontsize=6.8,
                color=MUTED,
            )

        axis.axvline(1.75, color=GRID, lw=0.8, zorder=0)
        axis.set_xlim(-0.45, 3.95)
        axis.set_ylim(*ylim)
        axis.set_ylabel(ylabel, fontsize=8.5)
        axis.set_xticks(positions)
        axis.set_xticklabels([group[2] for group in groups], fontsize=7.0)
        axis.text(
            0.45,
            -0.25,
            "Flat, half-scale",
            transform=axis.get_xaxis_transform(),
            ha="center",
            va="top",
            fontsize=8.0,
            fontweight="bold",
            color=INK,
        )
        axis.text(
            3.05,
            -0.25,
            "Upright, full-scale",
            transform=axis.get_xaxis_transform(),
            ha="center",
            va="top",
            fontsize=8.0,
            fontweight="bold",
            color=INK,
        )
        axis.grid(axis="y", color=GRID, lw=0.65)
        axis.spines[["top", "right"]].set_visible(False)
        axis.tick_params(axis="y", labelsize=7.5)

    figure.savefig(output_path, dpi=600, bbox_inches="tight", facecolor="white")
    plt.close(figure)


def write_summary(specimens: list[Specimen], output_path: Path) -> None:
    fields = [
        "specimen_id",
        "source_filename",
        "condition",
        "condition_label",
        "orientation",
        "orientation_label",
        "scale",
        "replicate",
        "primary_analysis",
        "analysis_status",
        "raw_point_count",
        "curve_plot_point_count",
        "gauge_width_mm",
        "gauge_thickness_mm",
        "gauge_area_mm2",
        "initial_grip_separation_mm",
        "peak_force_n",
        "peak_time_s",
        "crosshead_travel_at_peak_mm",
        "crosshead_strain_at_peak",
        "crosshead_strain_at_peak_percent",
        "release_time_s",
        "crosshead_travel_at_release_mm",
        "crosshead_strain_at_release",
        "crosshead_strain_at_release_percent",
        "crosshead_work_density_to_peak_mj_m3",
        "crosshead_work_density_to_release_mj_m3",
        "ultimate_tensile_strength_mpa",
    ]
    with output_path.open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for specimen in specimens:
            writer.writerow(
                {
                    "specimen_id": specimen.specimen_id,
                    "source_filename": specimen.source_filename,
                    "condition": specimen.condition,
                    "condition_label": specimen.condition_label,
                    "orientation": specimen.orientation,
                    "orientation_label": specimen.orientation_label,
                    "scale": specimen.scale,
                    "replicate": specimen.replicate,
                    "primary_analysis": "yes" if specimen.is_primary else "no",
                    "analysis_status": specimen.status,
                    "raw_point_count": len(specimen.time_s),
                    "curve_plot_point_count": len(specimen.curve_time_s),
                    "gauge_width_mm": f"{specimen.width_mm:.2f}",
                    "gauge_thickness_mm": f"{specimen.thickness_mm:.2f}",
                    "gauge_area_mm2": f"{specimen.area_mm2:.4f}",
                    "initial_grip_separation_mm": f"{specimen.grip_separation_mm:.1f}",
                    "peak_force_n": f"{specimen.peak_force_n:.1f}",
                    "peak_time_s": f"{specimen.peak_time_s:.3f}",
                    "crosshead_travel_at_peak_mm": f"{specimen.peak_travel_mm:.4f}",
                    "crosshead_strain_at_peak": f"{specimen.peak_crosshead_strain:.5f}",
                    "crosshead_strain_at_peak_percent": f"{specimen.peak_crosshead_strain * 100.0:.3f}",
                    "release_time_s": f"{specimen.release_time_s:.3f}",
                    "crosshead_travel_at_release_mm": f"{specimen.release_travel_mm:.4f}",
                    "crosshead_strain_at_release": f"{specimen.release_crosshead_strain:.5f}",
                    "crosshead_strain_at_release_percent": f"{specimen.release_crosshead_strain * 100.0:.3f}",
                    "crosshead_work_density_to_peak_mj_m3": f"{specimen.crosshead_work_density_to_peak_mj_m3:.5f}",
                    "crosshead_work_density_to_release_mj_m3": f"{specimen.crosshead_work_density_to_release_mj_m3:.5f}",
                    "ultimate_tensile_strength_mpa": f"{specimen.peak_uts_mpa:.3f}",
                }
            )


def write_group_summary(specimens: list[Specimen], output_path: Path) -> None:
    fields = [
        "orientation",
        "orientation_label",
        "condition",
        "condition_label",
        "n_primary",
        "uts_mean_mpa",
        "uts_sd_mpa",
        "uts_ci95_half_width_mpa",
        "crosshead_strain_at_peak_mean_percent",
        "crosshead_strain_at_peak_sd_percent",
        "crosshead_strain_at_peak_ci95_half_width_percent",
        "crosshead_work_density_to_peak_mean_mj_m3",
        "crosshead_work_density_to_peak_sd_mj_m3",
        "crosshead_work_density_to_peak_ci95_half_width_mj_m3",
    ]
    with output_path.open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for orientation, condition in (
            ("flat", "magenta"),
            ("flat", "cm_weave"),
            ("upright", "magenta"),
            ("upright", "cm_weave"),
        ):
            group = primary_group(specimens, orientation, condition)
            uts = [specimen.peak_uts_mpa for specimen in group]
            strain = [specimen.peak_crosshead_strain * 100.0 for specimen in group]
            work = [specimen.crosshead_work_density_to_peak_mj_m3 for specimen in group]
            writer.writerow(
                {
                    "orientation": orientation,
                    "orientation_label": group[0].orientation_label,
                    "condition": condition,
                    "condition_label": group[0].condition_label,
                    "n_primary": len(group),
                    "uts_mean_mpa": f"{mean(uts):.4f}",
                    "uts_sd_mpa": f"{sample_sd(uts):.4f}",
                    "uts_ci95_half_width_mpa": f"{ci95(uts):.4f}",
                    "crosshead_strain_at_peak_mean_percent": f"{mean(strain):.4f}",
                    "crosshead_strain_at_peak_sd_percent": f"{sample_sd(strain):.4f}",
                    "crosshead_strain_at_peak_ci95_half_width_percent": f"{ci95(strain):.4f}",
                    "crosshead_work_density_to_peak_mean_mj_m3": f"{mean(work):.5f}",
                    "crosshead_work_density_to_peak_sd_mj_m3": f"{sample_sd(work):.5f}",
                    "crosshead_work_density_to_peak_ci95_half_width_mj_m3": f"{ci95(work):.5f}",
                }
            )


def validate_results(specimens: list[Specimen]) -> None:
    """Guard the figures against accidental changes to input/inclusion logic."""
    expected = {
        ("flat", "magenta"): (5, 35.051),
        ("flat", "cm_weave"): (5, 36.447),
        ("upright", "magenta"): (3, 15.400),
        ("upright", "cm_weave"): (3, 15.916),
    }
    for key, (expected_count, expected_mean) in expected.items():
        group = primary_group(specimens, *key)
        if len(group) != expected_count:
            raise AssertionError(f"{key} has {len(group)} primary specimens; expected {expected_count}")
        observed_mean = mean(specimen.peak_uts_mpa for specimen in group)
        if not math.isclose(observed_mean, expected_mean, rel_tol=0.0, abs_tol=0.006):
            raise AssertionError(f"{key} UTS mean {observed_mean:.3f} differs from expected {expected_mean:.3f}")
    flat_retention = mean(specimen.peak_uts_mpa for specimen in primary_group(specimens, "flat", "cm_weave")) / mean(specimen.peak_uts_mpa for specimen in primary_group(specimens, "flat", "magenta")) * 100.0
    upright_retention = mean(specimen.peak_uts_mpa for specimen in primary_group(specimens, "upright", "cm_weave")) / mean(specimen.peak_uts_mpa for specimen in primary_group(specimens, "upright", "magenta")) * 100.0
    if not math.isclose(flat_retention, 103.982, rel_tol=0.0, abs_tol=0.02):
        raise AssertionError(f"Flat CM retention changed: {flat_retention:.3f}%")
    if not math.isclose(upright_retention, 103.354, rel_tol=0.0, abs_tol=0.02):
        raise AssertionError(f"Upright CM retention changed: {upright_retention:.3f}%")

    expected_crosshead = {
        ("flat", "magenta"): (6.000, 1.003),
        ("flat", "cm_weave"): (5.885, 1.018),
        ("upright", "magenta"): (3.410, 0.261),
        ("upright", "cm_weave"): (3.943, 0.302),
    }
    for key, (expected_strain, expected_work) in expected_crosshead.items():
        group = primary_group(specimens, *key)
        observed_strain = mean(specimen.peak_crosshead_strain * 100.0 for specimen in group)
        observed_work = mean(specimen.crosshead_work_density_to_peak_mj_m3 for specimen in group)
        if not math.isclose(observed_strain, expected_strain, rel_tol=0.0, abs_tol=0.002):
            raise AssertionError(
                f"{key} mean crosshead strain at UTS {observed_strain:.3f}% "
                f"differs from expected {expected_strain:.3f}%"
            )
        if not math.isclose(observed_work, expected_work, rel_tol=0.0, abs_tol=0.002):
            raise AssertionError(
                f"{key} mean work density to UTS {observed_work:.3f} MJ/m^3 "
                f"differs from expected {expected_work:.3f} MJ/m^3"
            )


def main() -> None:
    revision_dir = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--data-dir",
        type=Path,
        default=revision_dir / "data" / "tensile_raw",
    )
    parser.add_argument("--revision-dir", type=Path, default=revision_dir)
    args = parser.parse_args()

    paths = sorted(args.data_dir.glob("*.xlsx"))
    if len(paths) != 20:
        raise RuntimeError(f"Expected 20 Mark-10 exports in {args.data_dir}, found {len(paths)}")
    specimens = [parse_specimen(path) for path in paths]
    validate_results(specimens)

    supplementary_figures_dir = args.revision_dir / "supplement" / "figures"
    tables_dir = args.revision_dir / "supplement" / "tables"
    supplementary_figures_dir.mkdir(parents=True, exist_ok=True)
    tables_dir.mkdir(parents=True, exist_ok=True)
    plot_main_figure(specimens, supplementary_figures_dir / "supp_fig09_tensile_strength_screen.png")
    plot_crosshead_curves(specimens, supplementary_figures_dir / "supp_fig10_tensile_crosshead_curves.png")
    plot_crosshead_endpoints(
        specimens,
        supplementary_figures_dir / "supp_fig11_tensile_crosshead_endpoints.png",
    )
    write_summary(specimens, tables_dir / "tensile_screen_specimen_summary.csv")
    write_group_summary(specimens, tables_dir / "tensile_screen_group_summary.csv")

    for orientation, condition in (("flat", "magenta"), ("flat", "cm_weave"), ("upright", "magenta"), ("upright", "cm_weave")):
        group = primary_group(specimens, orientation, condition)
        values = [specimen.peak_uts_mpa for specimen in group]
        peak_strain = [specimen.peak_crosshead_strain * 100.0 for specimen in group]
        work = [specimen.crosshead_work_density_to_peak_mj_m3 for specimen in group]
        print(
            f"{orientation:7s} {condition:9s} n={len(values)} "
            f"UTS={mean(values):.3f}+/-{sample_sd(values):.3f} MPa; "
            f"crosshead strain at UTS={mean(peak_strain):.3f}+/-{sample_sd(peak_strain):.3f}%; "
            f"work to UTS={mean(work):.3f}+/-{sample_sd(work):.3f} MJ/m^3"
        )


if __name__ == "__main__":
    main()
